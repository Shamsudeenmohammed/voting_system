"""Tests for the elections app: lifecycle, four-eyes, configuration, registry.

Shared fixtures for other apps' tests live in :class:`BaseVotingTestCase`.
"""
from __future__ import annotations

from datetime import timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from organizations.models import Organization, OrganizationMembership

from .config import SelectionError, compute_config_hash, validate_selections
from .lifecycle import LifecycleError, request_approval, decide_approval, sync_election_state, transition
from .models import Candidate, Election, Position, Voter

User = get_user_model()


class BaseVotingTestCase(TestCase):
    """One organization with an admin, an official, an auditor and an election."""

    @classmethod
    def setUpTestData(cls):
        cls.org = Organization.objects.create(name="Uni", slug="uni")
        cls.other_org = Organization.objects.create(name="Other", slug="other")
        cls.admin = User.objects.create_user("admin", password="x", is_staff=True)
        cls.admin2 = User.objects.create_user("admin2", password="x", is_staff=True)
        cls.official = User.objects.create_user("official", password="x", is_staff=True)
        cls.auditor = User.objects.create_user("auditor", password="x", is_staff=True)
        cls.outsider = User.objects.create_user("outsider", password="x", is_staff=True)
        OrganizationMembership.objects.create(
            user=cls.admin, organization=cls.org, role="org_admin"
        )
        OrganizationMembership.objects.create(
            user=cls.admin2, organization=cls.org, role="org_admin"
        )
        OrganizationMembership.objects.create(
            user=cls.official, organization=cls.org, role="election_official"
        )
        OrganizationMembership.objects.create(
            user=cls.auditor, organization=cls.org, role="auditor"
        )
        # outsider belongs to the *other* organization only
        OrganizationMembership.objects.create(
            user=cls.outsider, organization=cls.other_org, role="org_admin"
        )

    def make_election(self, **kwargs) -> Election:
        now = timezone.now()
        defaults = dict(
            organization=self.org,
            title="Students' Union Election",
            start_at=now - timedelta(hours=1),
            end_at=now + timedelta(hours=3),
            created_by=self.official,
        )
        defaults.update(kwargs)
        election = Election.objects.create(**defaults)
        return election

    def make_ballot_structure(self, election, *, max_choices=1) -> tuple[Position, list[Candidate]]:
        position = Position.objects.create(
            election=election, title="President", max_choices=max_choices
        )
        candidates = [
            Candidate.objects.create(position=position, full_name="Alice A."),
            Candidate.objects.create(position=position, full_name="Bob B."),
        ]
        return position, candidates

    def make_voters(self, election, n=3) -> list[Voter]:
        return [
            Voter.objects.create(
                election=election, voter_id=f"V{i:03d}", full_name=f"Voter {i}"
            )
            for i in range(1, n + 1)
        ]

    def open_election(self, *, with_voters=3) -> tuple[Election, Position, list[Candidate]]:
        """Election in OPEN state with a complete ballot and voters."""
        election = self.make_election()
        position, candidates = self.make_ballot_structure(election)
        self.make_voters(election, with_voters)
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        election = transition(election, Election.STATE_CONFIG_REVIEW, actor=self.official)
        approval = request_approval(
            election, "approve_config", actor=self.admin, reason="Reviewed."
        )
        decide_approval(approval.pk, actor=self.admin2, approve=True)
        election.refresh_from_db()
        election = transition(election, Election.STATE_SCHEDULED, actor=self.official)
        election = transition(election, Election.STATE_OPEN, actor=self.official)
        return election, position, candidates


class LifecycleTests(BaseVotingTestCase):
    def test_only_allowed_edges(self):
        election = self.make_election()
        with self.assertRaises(LifecycleError) as ctx:
            transition(election, Election.STATE_OPEN, actor=self.official)
        self.assertEqual(ctx.exception.code, "illegal_transition")

    def test_draft_to_registration(self):
        election = self.make_election()
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        self.assertEqual(election.state, Election.STATE_REGISTRATION)

    def test_config_freeze_hashes_and_credential_issuance(self):
        election, position, candidates = self.open_election()
        self.assertTrue(election.config_hash)
        self.assertIsNotNone(election.config_frozen_at)
        # every active voter received a credential at freeze time
        from ballot.models import BallotAuthorization

        self.assertEqual(
            BallotAuthorization.objects.filter(election=election).count(), 3
        )

    def test_direct_approval_blocked(self):
        election = self.make_election()
        self.make_ballot_structure(election)
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        election = transition(election, Election.STATE_CONFIG_REVIEW, actor=self.official)
        with self.assertRaises(LifecycleError) as ctx:
            transition(election, Election.STATE_APPROVED, actor=self.admin)
        self.assertEqual(ctx.exception.code, "approval_required")

    def test_four_eyes_requester_cannot_decide(self):
        election = self.make_election()
        self.make_ballot_structure(election)
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        election = transition(election, Election.STATE_CONFIG_REVIEW, actor=self.official)
        approval = request_approval(
            election, "approve_config", actor=self.admin, reason="Looks good."
        )
        with self.assertRaises(LifecycleError) as ctx:
            decide_approval(approval.pk, actor=self.admin, approve=True)
        self.assertEqual(ctx.exception.code, "four_eyes")
        approval = decide_approval(approval.pk, actor=self.admin2, approve=True)
        election.refresh_from_db()
        self.assertEqual(approval.status, "approved")
        self.assertEqual(election.state, Election.STATE_APPROVED)

    def test_official_cannot_decide_approvals(self):
        election = self.make_election()
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        election = transition(election, Election.STATE_CONFIG_REVIEW, actor=self.official)
        approval = request_approval(
            election, "approve_config", actor=self.admin, reason="ok"
        )
        with self.assertRaises(LifecycleError) as ctx:
            decide_approval(approval.pk, actor=self.official, approve=True)
        self.assertEqual(ctx.exception.code, "forbidden")

    def test_approval_requires_reason(self):
        election = self.make_election()
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        election = transition(election, Election.STATE_CONFIG_REVIEW, actor=self.official)
        with self.assertRaises(LifecycleError) as ctx:
            request_approval(election, "approve_config", actor=self.admin, reason="  ")
        self.assertEqual(ctx.exception.code, "reason_required")

    def test_reopen_needs_four_eyes_and_counts(self):
        election, _, _ = self.open_election()
        election = transition(election, Election.STATE_CLOSED, actor=self.official)
        with self.assertRaises(LifecycleError):
            transition(election, Election.STATE_OPEN, actor=self.official)
        approval = request_approval(
            election, "reopen", actor=self.admin, reason="Station blackout."
        )
        decide_approval(approval.pk, actor=self.admin2, approve=True)
        election.refresh_from_db()
        self.assertEqual(election.state, Election.STATE_OPEN)
        self.assertEqual(election.reopened_count, 1)

    def test_decertify_supersedes_certified_snapshot(self):
        from ballot.services import cast_ballot, prepare_ballot
        from tally import services as tally_services
        from tally.models import TallySnapshot

        election, position, candidates = self.open_election()
        prep = prepare_ballot(election.voters.first(), election)
        cast_ballot(
            election=election,
            voter=election.voters.first(),
            submission_id=prep.submission_id,
            selections={str(position.pk): [str(candidates[0].pk)]},
        )
        election = transition(election, Election.STATE_CLOSED, actor=self.official)
        election = transition(election, Election.STATE_RECONCILIATION, actor=self.official)
        election = transition(election, Election.STATE_REVIEW, actor=self.official)
        election = transition(election, Election.STATE_CERTIFIED, actor=self.official)
        snapshot = tally_services.latest_snapshot(election)
        self.assertEqual(snapshot.status, TallySnapshot.STATUS_CERTIFIED)

        approval = request_approval(
            election, "decertify", actor=self.admin, reason="Recount requested."
        )
        decide_approval(approval.pk, actor=self.admin2, approve=True)
        election.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(election.state, Election.STATE_REVIEW)
        self.assertEqual(snapshot.status, TallySnapshot.STATUS_SUPERSEDED)

    def test_cannot_certify_without_reconciliation(self):
        election, _, _ = self.open_election()
        election = transition(election, Election.STATE_CLOSED, actor=self.official)
        # jumping over RECONCILIATION/REVIEW is not an allowed edge
        with self.assertRaises(LifecycleError) as ctx:
            transition(election, Election.STATE_CERTIFIED, actor=self.admin)
        self.assertEqual(ctx.exception.code, "illegal_transition")

    def test_schedule_driven_open_and_close(self):
        now = timezone.now()
        election = self.make_election(
            start_at=now - timedelta(minutes=5), end_at=now + timedelta(minutes=5)
        )
        self.make_ballot_structure(election)
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        election = transition(election, Election.STATE_CONFIG_REVIEW, actor=self.official)
        approval = request_approval(
            election, "approve_config", actor=self.admin, reason="ok"
        )
        decide_approval(approval.pk, actor=self.admin2, approve=True)
        election = transition(election, Election.STATE_SCHEDULED, actor=self.official)
        # start time has passed -> the system opens it
        election = sync_election_state(election)
        self.assertEqual(election.state, Election.STATE_OPEN)
        # ... and closes it once the end time passed
        Election.objects.filter(pk=election.pk).update(
            end_at=now - timedelta(minutes=1)
        )
        election.refresh_from_db()
        election = sync_election_state(election)
        self.assertEqual(election.state, Election.STATE_CLOSED)

    def test_state_transitions_are_audited(self):
        from audit.models import AuditEvent

        election = self.make_election()
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        self.assertTrue(
            AuditEvent.objects.filter(
                election=election, action="election_transition"
            ).exists()
        )


class ConfigTests(BaseVotingTestCase):
    def test_config_hash_is_deterministic(self):
        election = self.make_election()
        self.make_ballot_structure(election)
        self.assertEqual(compute_config_hash(election), compute_config_hash(election))

    def test_config_hash_changes_with_candidates(self):
        election = self.make_election()
        position, candidates = self.make_ballot_structure(election)
        before = compute_config_hash(election)
        Candidate.objects.create(position=position, full_name="Carol C.")
        self.assertNotEqual(before, compute_config_hash(election))

    def test_validate_selections_overvote(self):
        election = self.make_election()
        position, candidates = self.make_ballot_structure(election, max_choices=1)
        with self.assertRaises(SelectionError):
            validate_selections(
                election, {str(position.pk): [str(c.pk) for c in candidates]}
            )

    def test_validate_selections_unknown_candidate(self):
        election = self.make_election()
        position, _ = self.make_ballot_structure(election)
        with self.assertRaises(SelectionError):
            validate_selections(election, {str(position.pk): ["not-a-candidate"]})

    def test_blank_rejected_when_disallowed(self):
        election = self.make_election(allow_blank=False)
        position, _ = self.make_ballot_structure(election)
        with self.assertRaises(SelectionError):
            validate_selections(election, {str(position.pk): []})

    def test_canonical_selection_order(self):
        election = self.make_election()
        position, candidates = self.make_ballot_structure(election, max_choices=2)
        result = validate_selections(
            election, {str(position.pk): [str(candidates[1].pk), str(candidates[0].pk)]}
        )
        expected = sorted([str(candidates[0].pk), str(candidates[1].pk)])
        self.assertEqual(result[str(position.pk)], expected)


class RegistryTests(BaseVotingTestCase):
    def test_voters_only_during_registration(self):
        election = self.make_election()
        Voter.objects.create(election=election, voter_id="X1")
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        self.assertTrue(election.registration_open)
        election = transition(election, Election.STATE_CONFIG_REVIEW, actor=self.official)
        self.assertFalse(election.registration_open)

    def test_duplicate_voter_id_rejected(self):
        election = self.make_election()
        election = transition(election, Election.STATE_REGISTRATION, actor=self.official)
        Voter.objects.create(election=election, voter_id="DUP")
        from django.db import IntegrityError

        with self.assertRaises(IntegrityError):
            Voter.objects.create(election=election, voter_id="DUP")

    def test_blank_legacy_ids_allowed(self):
        election = self.make_election()
        Voter.objects.create(election=election, voter_id="")
        Voter.objects.create(election=election, voter_id="")
        self.assertEqual(election.voters.count(), 2)

    def test_eligibility_units(self):
        election = self.make_election(eligibility_units=["Faculty A"])
        v_ok = Voter.objects.create(
            election=election, voter_id="A1", organization_unit="Faculty A"
        )
        v_bad = Voter.objects.create(
            election=election, voter_id="B1", organization_unit="Faculty B"
        )
        self.assertTrue(v_ok.eligible_for(election))
        self.assertFalse(v_bad.eligible_for(election))


class RegistryImportTests(BaseVotingTestCase):
    """Downloadable register template and file (CSV/XLSX) upload."""

    def _messages(self, response) -> list[str]:
        from django.contrib.messages import get_messages

        return [str(m) for m in get_messages(response.wsgi_request)]

    def _registration_election(self) -> Election:
        election = self.make_election()
        return transition(election, Election.STATE_REGISTRATION, actor=self.official)

    # -- template download ------------------------------------------------

    def test_template_download_csv(self):
        election = self._registration_election()
        self.client.force_login(self.official)
        r = self.client.get(reverse("voter_registry_template", args=[election.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/csv", r["Content-Type"])
        self.assertIn("attachment", r["Content-Disposition"])
        lines = r.content.decode("utf-8").strip().splitlines()
        self.assertEqual(lines[0], "voter_id,full_name,email,unit")
        self.assertIn("V001,Ada Lovelace,ada@example.org,Science Faculty", lines)

    def test_template_download_xlsx(self):
        election = self._registration_election()
        self.client.force_login(self.official)
        r = self.client.get(
            reverse("voter_registry_template", args=[election.pk]) + "?format=xlsx"
        )
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])

        import io

        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("voter_id", "full_name", "email", "unit"))
        self.assertEqual(rows[1][0], "V001")

    def test_template_download_scoped(self):
        election = self._registration_election()
        self.client.force_login(self.outsider)
        r = self.client.get(reverse("voter_registry_template", args=[election.pk]))
        self.assertEqual(r.status_code, 403)

    def test_template_round_trip_import(self):
        """Download the template, fill it, upload it back."""
        import io

        import openpyxl

        from .voter_import import template_xlsx_bytes

        election = self._registration_election()
        wb = openpyxl.load_workbook(io.BytesIO(template_xlsx_bytes()))
        ws = wb.active
        ws.append(("V900", "Test User", "test@example.org", "Law Faculty"))
        buffer = io.BytesIO()
        wb.save(buffer)

        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_login(self.official)
        r = self.client.post(
            reverse("voter_file_add", args=[election.pk]),
            {
                "file": SimpleUploadedFile(
                    "register.xlsx",
                    buffer.getvalue(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument"
                        ".spreadsheetml.sheet"
                    ),
                )
            },
        )
        self.assertEqual(r.status_code, 302, self._messages(r))
        self.assertEqual(election.voters.count(), 4)  # 3 template rows + 1 added

    # -- file upload ------------------------------------------------------

    def _upload_csv(self, election: Election, content: str, **extra):
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_login(self.official)
        return self.client.post(
            reverse("voter_file_add", args=[election.pk]),
            {
                "file": SimpleUploadedFile(
                    "register.csv", content.encode("utf-8"), content_type="text/csv"
                ),
                **extra,
            },
        )

    def test_file_upload_registers_voters(self):
        election = self._registration_election()
        csv = (
            "voter_id,full_name,email,unit\n"
            "F001,Ada Lovelace,ada@example.org,Science Faculty\n"
            "F002,Grace Hopper,,Engineering Faculty\n"
        )
        r = self._upload_csv(election, csv)
        self.assertEqual(r.status_code, 302, self._messages(r))
        self.assertEqual(election.voters.count(), 2)
        voter = election.voters.get(voter_id="F002")
        self.assertEqual(voter.organization_unit, "Engineering Faculty")
        self.assertTrue(
            any("2 voters registered from file" in m for m in self._messages(r)),
            self._messages(r),
        )

    def test_file_upload_skips_existing_ids(self):
        election = self._registration_election()
        Voter.objects.create(election=election, voter_id="F001", full_name="Old Row")
        csv = "voter_id,full_name\nF001,Ada Lovelace\nF002,Grace Hopper\n"
        r = self._upload_csv(election, csv)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(election.voters.count(), 2)
        self.assertTrue(
            any("skipped" in m for m in self._messages(r)),
            self._messages(r),
        )
        self.assertEqual(
            election.voters.get(voter_id="F001").full_name, "Old Row"
        )

    def test_file_upload_duplicate_in_file_rejected(self):
        election = self._registration_election()
        csv = "voter_id,full_name\nF001,Ada Lovelace\nF001,Duplicate\n"
        r = self._upload_csv(election, csv)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(election.voters.count(), 0)
        self.assertTrue(
            any("duplicate entry in file" in m for m in self._messages(r)),
            self._messages(r),
        )

    def test_file_upload_requires_id_or_name(self):
        election = self._registration_election()
        csv = "voter_id,full_name,email\n,,"
        r = self._upload_csv(election, csv)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(election.voters.count(), 0)

    def test_file_upload_rejects_unsupported_type(self):
        election = self._registration_election()
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_login(self.official)
        r = self.client.post(
            reverse("voter_file_add", args=[election.pk]),
            {"file": SimpleUploadedFile("register.pdf", b"%PDF-1.4")},
        )
        self.assertEqual(r.status_code, 302)
        self.assertEqual(election.voters.count(), 0)
        self.assertTrue(
            any("Unsupported file type" in m for m in self._messages(r)),
            self._messages(r),
        )

    def test_file_upload_requires_file(self):
        election = self._registration_election()
        self.client.force_login(self.official)
        r = self.client.post(reverse("voter_file_add", args=[election.pk]), {})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(election.voters.count(), 0)

    def test_registry_page_offers_template_and_file_input(self):
        election = self._registration_election()
        Voter.objects.create(election=election, voter_id="P001", full_name="Pre")
        self.client.force_login(self.official)
        r = self.client.get(reverse("election_voters", args=[election.pk]))
        self.assertEqual(r.status_code, 200)
        template_url = reverse("voter_registry_template", args=[election.pk])
        self.assertContains(r, template_url)  # download CSV link
        self.assertContains(r, f"{template_url}?format=xlsx")
        self.assertContains(r, 'type="file"')  # file chooser
        self.assertContains(r, "multipart/form-data")

    def test_file_upload_only_during_registration(self):
        election = self.make_election()  # DRAFT, not REGISTRATION
        r = self._upload_csv(
            election, "voter_id,full_name\nF001,Ada Lovelace\n"
        )
        self.assertEqual(r.status_code, 302)
        self.assertEqual(election.voters.count(), 0)
        self.assertTrue(
            any("REGISTRATION" in m for m in self._messages(r)),
            self._messages(r),
        )
