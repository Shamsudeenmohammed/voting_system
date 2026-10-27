"""Voter registry file import helpers.

Provides the downloadable register template (CSV/XLSX) and a parser that
turns an uploaded file into the same row dictionaries used by the pasted-CSV
bulk form, so both ingestion paths share one validation contract:

    voter_id, full_name (at least one required)
    email, unit (optional)
"""
from __future__ import annotations

import csv
import io

MAX_FILE_ROWS = 5000

TEMPLATE_COLUMNS = ("voter_id", "full_name", "email", "unit")
TEMPLATE_SAMPLE_ROWS = [
    ("V001", "Ada Lovelace", "ada@example.org", "Science Faculty"),
    ("V002", "Grace Hopper", "grace@example.org", "Engineering Faculty"),
    ("V003", "Kwame Mensah", "kwame@example.org", "Arts Faculty"),
]


def template_csv_bytes() -> bytes:
    """Serialize the example register template as CSV (UTF-8, BOM-free)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(TEMPLATE_COLUMNS)
    writer.writerows(TEMPLATE_SAMPLE_ROWS)
    return buffer.getvalue().encode("utf-8")


def template_xlsx_bytes() -> bytes:
    """Serialize the example register template as an XLSX workbook."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Voters"
    ws.append(list(TEMPLATE_COLUMNS))
    for row in TEMPLATE_SAMPLE_ROWS:
        ws.append(list(row))
    for column_cells in ws.columns:  # readable default widths
        width = max(len(str(c.value or "")) for c in column_cells) + 4
        ws.column_dimensions[column_cells[0].column_letter].width = min(width, 40)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _normalize_row(row: dict) -> dict:
    return {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}


def _row_to_registry(row: dict) -> dict:
    return {
        "voter_id": row.get("voter_id", ""),
        "full_name": row.get("full_name", ""),
        "email": row.get("email") or None,
        "organization_unit": row.get("unit", ""),
    }


def _iter_rows(upload) -> "list[tuple[int, dict]]":
    """Yield ``(line_no, normalized_row)`` pairs from a CSV/TXT/XLSX upload."""
    name = (upload.name or "").lower()
    if name.endswith((".csv", ".txt")):
        raw = upload.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("The file is empty; download the template for the expected format.")
        return [
            (line_no, _normalize_row(row))
            for line_no, row in enumerate(reader, start=2)
        ]
    if name.endswith((".xlsx", ".xls")):
        import openpyxl

        workbook = openpyxl.load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValueError("The file is empty; download the template for the expected format.")
        keys = [str(h).strip().lower() if h is not None else "" for h in header]
        parsed = []
        for line_no, values in enumerate(rows, start=2):
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            parsed.append(
                (
                    line_no,
                    {k: ("" if v is None else str(v).strip()) for k, v in zip(keys, values)},
                )
            )
        return parsed
    raise ValueError("Unsupported file type. Use CSV, TXT or XLSX.")


def rows_from_upload(upload) -> "list[dict]":
    """Parse an uploaded register file into registry row dictionaries.

    Raises ``ValueError`` with a user-readable message for unsupported,
    empty, oversized or malformed files (matching the strict behaviour of
    the pasted-CSV bulk form). Blank lines are skipped.
    """
    rows: list[dict] = []
    seen: set[str] = set()
    for line_no, row in _iter_rows(upload):
        registry_row = _row_to_registry(row)
        voter_id = registry_row["voter_id"]
        full_name = registry_row["full_name"]
        if not voter_id and not full_name:
            # Tolerate stray blank lines / trailing newlines.
            if not any(registry_row.values()):
                continue
            raise ValueError(
                f"Line {line_no}: a voter_id or full_name is required."
            )
        key = voter_id.lower() or f"__name__:{full_name.lower()}"
        if key in seen:
            raise ValueError(
                f"Line {line_no}: duplicate entry in file: {voter_id or full_name}."
            )
        seen.add(key)
        rows.append(registry_row)
        if len(rows) > MAX_FILE_ROWS:
            raise ValueError(
                f"File import is limited to {MAX_FILE_ROWS:,} rows per upload."
            )
    if not rows:
        raise ValueError("No voter rows found in the file.")
    return rows
